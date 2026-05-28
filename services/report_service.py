import re
from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

from services.config_service import BASE_DIR, load_settings


def get_report_directory(report_type: str) -> Path:
    settings = load_settings()
    report_paths = settings["paths"]["reports"]

    if report_type not in report_paths:
        raise ValueError(f"Unknown report type: {report_type}")

    report_dir = BASE_DIR / report_paths[report_type]

    if report_dir.exists() and not report_dir.is_dir():
        raise FileExistsError(
            f"{report_dir} exists but is a file. Delete it and create it as a folder."
        )

    report_dir.mkdir(parents=True, exist_ok=True)
    return report_dir


def save_markdown_report(report_type: str, title: str, content: str) -> Path:
    report_dir = get_report_directory(report_type)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    safe_title = title.lower().replace(" ", "_").replace("/", "-")
    filename = f"{timestamp}_{safe_title}.md"

    file_path = report_dir / filename

    with open(file_path, "w", encoding="utf-8") as file:
        file.write(content)

    return file_path


def list_reports(report_type: str) -> list[Path]:
    report_dir = get_report_directory(report_type)
    return sorted(report_dir.glob("*.md"), reverse=True)


def read_report(file_path: Path) -> str:
    with open(file_path, "r", encoding="utf-8") as file:
        return file.read()


def clean_markdown_text(value: str) -> str:
    value = value.strip()
    value = value.replace("**", "")
    value = value.replace("__", "")
    value = value.replace("`", "")
    return value


def is_markdown_table_row(line: str) -> bool:
    return line.strip().startswith("|") and line.strip().endswith("|")


def is_markdown_table_separator(line: str) -> bool:
    stripped = line.strip()

    if not is_markdown_table_row(stripped):
        return False

    cells = [cell.strip() for cell in stripped.strip("|").split("|")]

    return all(
        set(cell.replace(":", "")) <= {"-"} and "-" in cell
        for cell in cells
    )


def parse_markdown_table(table_lines: list[str]) -> list[list[str]]:
    rows = []

    for line in table_lines:
        if is_markdown_table_separator(line):
            continue

        cells = [
            clean_markdown_text(cell.strip())
            for cell in line.strip().strip("|").split("|")
        ]

        rows.append(cells)

    return rows


def add_markdown_table_to_doc(document: Document, table_lines: list[str]) -> None:
    rows = parse_markdown_table(table_lines)

    if not rows:
        return

    column_count = max(len(row) for row in rows)

    table = document.add_table(rows=0, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for row_index, row_data in enumerate(rows):
        row = table.add_row()

        for col_index in range(column_count):
            cell_text = row_data[col_index] if col_index < len(row_data) else ""
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)

            run = paragraph.add_run(cell_text)
            run.font.size = Pt(8.5)

            if row_index == 0:
                run.bold = True

    document.add_paragraph()


def add_title_page(document: Document, title: str) -> None:
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_para.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(20)

    subtitle_para = document.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle_run = subtitle_para.add_run("Portfolio Intelligence Terminal")
    subtitle_run.font.size = Pt(12)

    date_para = document.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    date_run = date_para.add_run(datetime.now().strftime("%d %B %Y"))
    date_run.font.size = Pt(10)

    document.add_paragraph()
    document.add_paragraph()


def markdown_to_docx(markdown_text: str, output_filename: str) -> Path:
    export_dir = BASE_DIR / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)

    safe_filename = output_filename.lower()
    safe_filename = safe_filename.replace(".md", "")
    safe_filename = re.sub(r"[^a-z0-9_\-]+", "_", safe_filename)
    output_path = export_dir / f"{safe_filename}.docx"

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.bold = True

    report_title = output_filename.replace(".md", "").replace("_", " ").title()

    for line in markdown_text.splitlines():
        if line.startswith("# "):
            report_title = clean_markdown_text(line.replace("# ", "", 1))
            break

    add_title_page(document, report_title)

    lines = markdown_text.splitlines()
    table_buffer = []

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            if table_buffer:
                add_markdown_table_to_doc(document, table_buffer)
                table_buffer = []
            continue

        if is_markdown_table_row(line):
            table_buffer.append(line)
            continue

        if table_buffer:
            add_markdown_table_to_doc(document, table_buffer)
            table_buffer = []

        if line.startswith("# "):
            document.add_heading(
                clean_markdown_text(line.replace("# ", "", 1)),
                level=1,
            )

        elif line.startswith("## "):
            document.add_heading(
                clean_markdown_text(line.replace("## ", "", 1)),
                level=2,
            )

        elif line.startswith("### "):
            document.add_heading(
                clean_markdown_text(line.replace("### ", "", 1)),
                level=3,
            )

        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_markdown_text(line.replace("- ", "", 1)))
            run.font.size = Pt(10)

        elif re.match(r"^\d+\.\s+", line):
            clean_line = re.sub(r"^\d+\.\s+", "", line)
            paragraph = document.add_paragraph(style="List Number")
            run = paragraph.add_run(clean_markdown_text(clean_line))
            run.font.size = Pt(10)

        else:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(6)

            bold_match = re.match(r"^\*\*(.+?)\*\*(.*)$", line)

            if bold_match:
                bold_part = clean_markdown_text(bold_match.group(1))
                rest = clean_markdown_text(bold_match.group(2))

                bold_run = paragraph.add_run(bold_part)
                bold_run.bold = True
                bold_run.font.size = Pt(10)

                if rest:
                    rest_run = paragraph.add_run(rest)
                    rest_run.font.size = Pt(10)
            else:
                run = paragraph.add_run(clean_markdown_text(line))
                run.font.size = Pt(10)

    if table_buffer:
        add_markdown_table_to_doc(document, table_buffer)

    footer = document.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("Generated by Portfolio Intelligence Terminal")
    footer_run.font.size = Pt(8)

    document.save(output_path)
    return output_path

def export_research_pack_to_docx(include_deep_digger: bool = True) -> Path:
    """
    Export the latest reports from all agents into one management-ready Word document.
    """
    export_dir = BASE_DIR / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    output_path = export_dir / f"{timestamp}_research_pack.docx"

    report_sections = [
        {
            "title": "Morning Analyst",
            "report_type": "morning_briefs",
        },
        {
            "title": "Thesis Tracker",
            "report_type": "thesis_reviews",
        },
        {
            "title": "Event Calendar",
            "report_type": "weekly_calendars",
        },
        {
            "title": "Market Curator",
            "report_type": "market_curator",
        },
    ]

    if include_deep_digger:
        report_sections.append(
            {
                "title": "Deep Digger",
                "report_type": "deep_research",
            }
        )

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.bold = True

    add_title_page(document, "Portfolio Intelligence Research Pack")

    intro = document.add_paragraph()
    intro.add_run(
        "This research pack combines the latest available outputs from the Portfolio Intelligence Terminal agents."
    )

    generated = document.add_paragraph()
    generated.add_run(f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}")

    document.add_page_break()

    included_count = 0

    for section_config in report_sections:
        section_title = section_config["title"]
        report_type = section_config["report_type"]

        reports = list_reports(report_type)

        if not reports:
            document.add_heading(section_title, level=1)
            document.add_paragraph("No report available for this section.")
            document.add_page_break()
            continue

        latest_report = reports[0]
        report_text = read_report(latest_report)

        document.add_heading(section_title, level=1)

        meta_para = document.add_paragraph()
        meta_run = meta_para.add_run(f"Source file: {latest_report.name}")
        meta_run.italic = True
        meta_run.font.size = Pt(8)

        lines = report_text.splitlines()
        table_buffer = []

        for raw_line in lines:
            line = raw_line.strip()

            if not line:
                if table_buffer:
                    add_markdown_table_to_doc(document, table_buffer)
                    table_buffer = []
                continue

            if is_markdown_table_row(line):
                table_buffer.append(line)
                continue

            if table_buffer:
                add_markdown_table_to_doc(document, table_buffer)
                table_buffer = []

            if line.startswith("# "):
                document.add_heading(
                    clean_markdown_text(line.replace("# ", "", 1)),
                    level=2,
                )

            elif line.startswith("## "):
                document.add_heading(
                    clean_markdown_text(line.replace("## ", "", 1)),
                    level=3,
                )

            elif line.startswith("### "):
                paragraph = document.add_paragraph()
                run = paragraph.add_run(clean_markdown_text(line.replace("### ", "", 1)))
                run.bold = True
                run.font.size = Pt(10)

            elif line.startswith("- "):
                paragraph = document.add_paragraph(style="List Bullet")
                run = paragraph.add_run(clean_markdown_text(line.replace("- ", "", 1)))
                run.font.size = Pt(10)

            elif re.match(r"^\d+\.\s+", line):
                clean_line = re.sub(r"^\d+\.\s+", "", line)
                paragraph = document.add_paragraph(style="List Number")
                run = paragraph.add_run(clean_markdown_text(clean_line))
                run.font.size = Pt(10)

            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(6)

                bold_match = re.match(r"^\*\*(.+?)\*\*(.*)$", line)

                if bold_match:
                    bold_part = clean_markdown_text(bold_match.group(1))
                    rest = clean_markdown_text(bold_match.group(2))

                    bold_run = paragraph.add_run(bold_part)
                    bold_run.bold = True
                    bold_run.font.size = Pt(10)

                    if rest:
                        rest_run = paragraph.add_run(rest)
                        rest_run.font.size = Pt(10)
                else:
                    run = paragraph.add_run(clean_markdown_text(line))
                    run.font.size = Pt(10)

        if table_buffer:
            add_markdown_table_to_doc(document, table_buffer)

        included_count += 1
        document.add_page_break()

    footer = document.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("Generated by Portfolio Intelligence Terminal")
    footer_run.font.size = Pt(8)

    document.save(output_path)
    return output_path

def export_portfolio_risk_report_to_docx(
    risk_df,
    positions: list[dict],
) -> Path:
    """
    Export the Portfolio Risk Dashboard into a formal Word report.
    """
    export_dir = BASE_DIR / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    output_path = export_dir / f"{timestamp}_portfolio_risk_report.docx"

    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Aptos"
        styles[style_name].font.bold = True

    add_title_page(document, "Portfolio Risk Report")

    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = generated.add_run(
        f"Generated: {datetime.now().strftime('%d %B %Y %H:%M')}"
    )
    run.font.size = Pt(10)

    document.add_page_break()

    total_positions = len(risk_df)
    total_weight = risk_df["Weight %"].sum() if not risk_df.empty else 0
    high_risk_count = len(risk_df[risk_df["Risk Band"] == "High"])
    medium_risk_count = len(risk_df[risk_df["Risk Band"] == "Medium"])
    average_completeness = (
        risk_df["Thesis Completeness %"].mean() if not risk_df.empty else 0
    )

    if not risk_df.empty:
        largest_position = risk_df.sort_values("Weight %", ascending=False).iloc[0]
        largest_position_text = (
            f"{largest_position['Ticker']} "
            f"({largest_position['Weight %']:.1f}%)"
        )
    else:
        largest_position_text = "N/A"

    document.add_heading("1. Executive Summary", level=1)

    summary_items = [
        f"Total portfolio positions reviewed: {total_positions}",
        f"Total mapped portfolio weight: {total_weight:.1f}%",
        f"Largest position: {largest_position_text}",
        f"High-risk positions identified: {high_risk_count}",
        f"Medium-risk positions identified: {medium_risk_count}",
        f"Average thesis completeness score: {average_completeness:.0f}%",
    ]

    for item in summary_items:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("2. Portfolio Risk Dashboard", level=1)

    dashboard_columns = [
        "Ticker",
        "Company",
        "Sector",
        "Weight %",
        "News Priority",
        "Kill Criteria",
        "Thesis Completeness %",
        "Risk Score",
        "Risk Band",
    ]

    table = document.add_table(rows=1, cols=len(dashboard_columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    header_cells = table.rows[0].cells
    for index, column in enumerate(dashboard_columns):
        run = header_cells[index].paragraphs[0].add_run(column)
        run.bold = True
        run.font.size = Pt(8)

    sorted_df = risk_df.sort_values(
        ["Risk Score", "Weight %"],
        ascending=[False, False],
    )

    for _, row in sorted_df.iterrows():
        cells = table.add_row().cells

        for index, column in enumerate(dashboard_columns):
            value = row.get(column, "")

            if column == "Weight %":
                value = f"{float(value):.1f}%"

            cell_run = cells[index].paragraphs[0].add_run(str(value))
            cell_run.font.size = Pt(8)

    document.add_paragraph()

    document.add_heading("3. Sector Exposure", level=1)

    sector_df = (
        risk_df.groupby("Sector", as_index=False)["Weight %"]
        .sum()
        .sort_values("Weight %", ascending=False)
    )

    sector_table = document.add_table(rows=1, cols=2)
    sector_table.style = "Table Grid"

    sector_table.rows[0].cells[0].text = "Sector"
    sector_table.rows[0].cells[1].text = "Weight %"

    for cell in sector_table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(8)

    for _, row in sector_df.iterrows():
        cells = sector_table.add_row().cells
        cells[0].text = str(row["Sector"])
        cells[1].text = f"{row['Weight %']:.1f}%"

    document.add_heading("4. High and Medium Risk Positions", level=1)

    review_df = risk_df[risk_df["Risk Band"].isin(["High", "Medium"])].sort_values(
        ["Risk Score", "Weight %"],
        ascending=[False, False],
    )

    if review_df.empty:
        document.add_paragraph("No high or medium risk positions identified.")
    else:
        for _, row in review_df.iterrows():
            document.add_heading(
                f"{row['Ticker']} — {row['Risk Band']} Risk",
                level=2,
            )

            details = [
                f"Company: {row['Company']}",
                f"Sector: {row['Sector']}",
                f"Weight: {row['Weight %']:.1f}%",
                f"News priority: {row['News Priority']}",
                f"Kill criteria count: {row['Kill Criteria']}",
                f"Thesis completeness: {row['Thesis Completeness %']}%",
                f"Risk score: {row['Risk Score']}",
            ]

            for item in details:
                document.add_paragraph(item, style="List Bullet")

            position = next(
                (
                    item for item in positions
                    if item.get("ticker") == row["Ticker"]
                ),
                None,
            )

            if position:
                thesis = position.get("investment_thesis", {})
                kill_items = thesis.get("kill_criteria", [])

                document.add_paragraph("Kill Criteria:")

                if kill_items:
                    for item in kill_items:
                        document.add_paragraph(item, style="List Bullet")
                else:
                    document.add_paragraph("No kill criteria recorded.")

    document.add_heading("5. Thesis Completeness Review", level=1)

    incomplete_df = risk_df.sort_values(
        "Thesis Completeness %",
        ascending=True,
    )

    for _, row in incomplete_df.iterrows():
        document.add_paragraph(
            f"{row['Ticker']}: {row['Thesis Completeness %']}% complete",
            style="List Bullet",
        )

    document.add_heading("6. Risk Methodology", level=1)

    methodology = [
        "This report is based on the information stored in the portfolio YAML files.",
        "The risk score considers position weight, monitoring priority, number of kill criteria, and thesis completeness.",
        "This is not a market-risk VaR model.",
        "It is a portfolio governance and thesis-risk dashboard designed to highlight areas requiring review.",
    ]

    for item in methodology:
        document.add_paragraph(item, style="List Bullet")

    footer = document.sections[0].footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_para.add_run("Generated by Portfolio Intelligence Terminal")
    footer_run.font.size = Pt(8)

    document.save(output_path)
    return output_path

def search_reports_for_text(report_type: str, search_text: str, limit: int = 5) -> list[dict]:
    """
    Search saved markdown reports for a ticker/company mention.
    """
    search_text = search_text.lower().strip()

    if not search_text:
        return []

    matches = []

    try:
        reports = list_reports(report_type)
    except Exception:
        return []

    for report_path in reports:
        try:
            text = read_report(report_path)
        except Exception:
            continue

        if search_text in text.lower() or search_text in report_path.name.lower():
            preview = text[:800] + "..." if len(text) > 800 else text

            matches.append(
                {
                    "file_name": report_path.name,
                    "path": report_path,
                    "preview": preview,
                    "modified": report_path.stat().st_mtime,
                }
            )

        if len(matches) >= limit:
            break

    return matches